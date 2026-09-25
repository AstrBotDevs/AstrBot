from __future__ import annotations

import io

import pytest

from astrbot.core.knowledge_base.parsers.markitdown_parser import MarkitdownParser
from astrbot.core.knowledge_base.parsers.util import select_parser


def _make_pptx_bytes() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "KB PPTX Slide Title"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_select_parser_supports_pptx():
    parser = await select_parser(".pptx")

    assert isinstance(parser, MarkitdownParser)


@pytest.mark.asyncio
async def test_select_parser_rejects_legacy_ppt():
    with pytest.raises(ValueError, match="暂时不支持的文件格式"):
        await select_parser(".ppt")


@pytest.mark.asyncio
async def test_pptx_parser_extracts_slide_text():
    pytest.importorskip("pptx")

    result = await MarkitdownParser().parse(_make_pptx_bytes(), "slides.pptx")

    assert result.media == []
    assert "KB PPTX Slide Title" in result.text
